"""Build the Falcon Tax Organizer workbook (tax year 2026).

Clients fill in the yellow cells and upload the file on the refund estimator
page. The last tab, "Data", turns every answer into the estimator's field
names with formulas; js/organizer-import.js reads that tab.

Run:  python3 tax-prep/tools/build_organizer.py [output.xlsx] [--sample]
Then recalculate (LibreOffice) so formula results are cached in the file.
"""
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
NAVY = "1D3557"
GOLD = "C8963E"
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")     # cells the client fills in
EXAMPLE_FILL = PatternFill("solid", fgColor="F2F2F2")
TOTAL_FILL = PatternFill("solid", fgColor="E7EEF6")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="C9D1DC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '$#,##0.00;($#,##0.00);"-"'
YEAR = "0"

STATUSES = [("Single", "single"), ("Married filing jointly", "mfj"),
            ("Married filing separately", "mfs"), ("Head of household", "hoh")]
STATES = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas", "CA": "California", "CO": "Colorado",
    "CT": "Connecticut", "DE": "Delaware", "DC": "District of Columbia", "FL": "Florida", "GA": "Georgia",
    "HI": "Hawaii", "ID": "Idaho", "IL": "Illinois", "IN": "Indiana", "IA": "Iowa", "KS": "Kansas",
    "KY": "Kentucky", "LA": "Louisiana", "ME": "Maine", "MD": "Maryland", "MA": "Massachusetts",
    "MI": "Michigan", "MN": "Minnesota", "MS": "Mississippi", "MO": "Missouri", "MT": "Montana",
    "NE": "Nebraska", "NV": "Nevada", "NH": "New Hampshire", "NJ": "New Jersey", "NM": "New Mexico",
    "NY": "New York", "NC": "North Carolina", "ND": "North Dakota", "OH": "Ohio", "OK": "Oklahoma",
    "OR": "Oregon", "PA": "Pennsylvania", "RI": "Rhode Island", "SC": "South Carolina", "SD": "South Dakota",
    "TN": "Tennessee", "TX": "Texas", "UT": "Utah", "VT": "Vermont", "VA": "Virginia", "WA": "Washington",
    "WV": "West Virginia", "WI": "Wisconsin", "WY": "Wyoming",
}

# (label, Schedule C line, deductible share, note)
EXPENSES = [
    ("GENERAL BUSINESS", None, None, None),
    ("Advertising & marketing", "8", 1, None),
    ("Bank, payment & platform fees", "10", 1, None),
    ("Contract labor (1099 workers)", "11", 1, "Send 1099-NEC to anyone paid $2,000+ in 2026."),
    ("Business insurance", "15", 1, None),
    ("Interest on business loans", "16b", 1, None),
    ("Legal & professional fees", "17", 1, None),
    ("Office expense & supplies", "18", 1, None),
    ("Rent (office, equipment)", "20b", 1, None),
    ("Taxes & licenses", "23", 1, None),
    ("Travel (overnight trips)", "24a", 1, None),
    ("Business meals (enter full amount)", "24b", 0.5, "Only 50% is deductible. Truck drivers: use the Trips tab instead. Rideshare and delivery drivers usually can't deduct meals."),
    ("Phone & internet (business share)", "25", 1, "Business-use percentage only."),
    ("Software & subscriptions", "27a", 1, None),
    ("Materials & supplies used for jobs", "22", 1, "Parts and materials you buy for customers' jobs."),
    ("Tools & equipment (under $2,500 each)", "22", 1, "Larger items may need depreciation."),
    ("Equipment rental", "20a", 1, None),
    ("Repairs (business equipment, non-vehicle)", "21", 1, None),
    ("Uniforms & safety gear", "27a", 1, "Only clothing not suitable for everyday wear."),
    ("Postage & shipping", "18", 1, None),
    ("Utilities (business location)", "25", 1, None),
    ("Training, courses & licenses for your trade", "27a", 1, None),
    ("Business gifts (up to $25 per person)", "27a", 1, None),
    ("Home office (we calculate the deductible part)", "30", 0, "Enter total rent/utilities for the home; we figure the business share (Form 8829)."),
    ("RIDESHARE & DELIVERY", None, None, None),
    ("Uber/Lyft/DoorDash service & booking fees", "10", 1, "From the annual tax summary in the driver app: fees taken out before you were paid."),
    ("Tolls & parking for business trips", "9", 1, "Deductible with either car method."),
    ("Passenger supplies (water, snacks, chargers)", "22", 1, None),
    ("Phone mount, cables, dash cam & car accessories", "22", 1, None),
    ("Delivery bags & gear", "22", 1, None),
    ("Airport & city rideshare permits", "23", 1, None),
    ("TRUCKING", None, None, None),
    ("Fuel (diesel)", "9", 1, None),
    ("DEF fluid", "9", 1, None),
    ("IFTA fuel tax paid", "23", 1, None),
    ("IFTA license & decals", "23", 1, None),
    ("Form 2290 heavy vehicle use tax", "23", 1, None),
    ("IRP plates & registration", "23", 1, None),
    ("UCR registration", "23", 1, None),
    ("Permits (trip, fuel, oversize/overweight)", "23", 1, None),
    ("State road taxes (NY HUT, KYU, NM WDT, OR)", "23", 1, None),
    ("Truck insurance (liability, cargo, physical damage)", "15", 1, None),
    ("Truck or trailer lease / rent", "20a", 1, None),
    ("Truck loan interest", "16b", 1, None),
    ("Tolls, scales & parking", "27a", 1, None),
    ("Load boards (DAT, Truckstop, 123Loadboard)", "27a", 1, None),
    ("ELD subscription", "27a", 1, None),
    ("Dispatch & factoring fees", "10", 1, None),
    ("Driver wages (W-2)", "26", 1, None),
    ("Contract drivers (1099)", "11", 1, None),
    ("Lumper fees", "27a", 1, None),
    ("Drug tests, DOT physicals & compliance", "27a", 1, None),
    ("Electronics (tablet, dash cam, GPS, ELD device)", "22", 1, "Items over $2,500 each may need depreciation."),
    ("Software purchase, one-time", "27a", 1, "Items over $2,500 each may need depreciation."),
    ("REPAIRS & MAINTENANCE (itemize)", None, None, None),
    ("Preventive maintenance (PM) service", "21", 1, None),
    ("Oil & filter change", "21", 1, None),
    ("Tires (new or retread)", "21", 1, None),
    ("Tire repair & alignment", "21", 1, None),
    ("Brakes", "21", 1, None),
    ("Engine repair", "21", 1, None),
    ("Transmission & clutch", "21", 1, None),
    ("DPF / emissions system", "21", 1, None),
    ("Electrical & batteries", "21", 1, None),
    ("Suspension, steering & air bags", "21", 1, None),
    ("A/C, heating & APU", "21", 1, None),
    ("Reefer unit service", "21", 1, None),
    ("Trailer repair", "21", 1, None),
    ("Body work, glass & paint", "21", 1, None),
    ("DOT annual inspection", "21", 1, None),
    ("Roadside service & towing", "21", 1, None),
    ("Truck wash", "21", 1, None),
    ("Parts & shop labor (other)", "21", 1, None),
    ("OTHER EXPENSES (describe)", None, None, None),
    ("", "27a", 1, None),
    ("", "27a", 1, None),
    ("", "27a", 1, None),
    ("", "27a", 1, None),
    ("", "27a", 1, None),
]

RIDESHARE_GROUP = "RIDESHARE & DELIVERY"

# Business car (rideshare, delivery or any business car, not semi-trucks).
# Miles are split at June 30: the IRS rate was 72.5 cents through June and
# 76 cents from July 1, 2026 (Notice 2026-10, Announcement 2026-11).
# The estimate compares standard mileage with actual costs x business share
# and uses the larger. Field names: car_<key>.
CAR_MILES = [
    ("miles_h1", "Business miles, Jan 1 - Jun 30, 2026"),
    ("miles_h2", "Business miles, Jul 1 - Dec 31, 2026"),
    ("miles_total", "ALL miles driven in 2026 (odometer Dec 31 minus Jan 1)"),
]
CAR_COSTS = [
    ("gas", "Gas or charging"),
    ("rent", "Car rental or lease payments (e.g., rental through Uber)"),
    ("insurance", "Car insurance (including rideshare coverage)"),
    ("repairs", "Repairs & maintenance (oil, tires, brakes)"),
    ("wash", "Car wash & detailing"),
    ("registration", "Registration, plates & inspection"),
    ("interest", "Car loan interest"),
    ("other", "Other car costs"),
]
MILEAGE_RATES = (0.725, 0.76)  # Jan-Jun, Jul-Dec 2026

CHECKLIST = [
    ("hasInvestments", "Did you sell stocks or crypto, or receive dividends?"),
    ("mortgage", "Did you pay mortgage interest on your home?"),
    ("studentLoan", "Did you pay student loan interest?"),
    ("college", "Did you pay college tuition for yourself or a dependent?"),
    ("marketplace", "Did anyone on your return have Marketplace health insurance (Form 1095-A)?"),
    ("noSsn", "Does anyone on the return NOT have a Social Security number?"),
    ("irsLetter", "Did you receive a letter from the IRS or the state?"),
]


def font(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("size", 10)
    return Font(**kw)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = font(size=16, bold=True, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = font(size=10, italic=True, color="5A6475")
    ws.sheet_view.showGridLines = False


def header(ws, row, labels, widths=None):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=i + 1, value=text)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w


def inp(c, fmt=None):
    c.fill = INPUT_FILL
    c.border = BOX
    c.font = font(color="0000FF")
    c.protection = Protection(locked=False)
    if fmt:
        c.number_format = fmt


def example(ws, row, values, fmts=None):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=i + 1, value=v)
        c.fill = EXAMPLE_FILL
        c.font = font(italic=True, color="7A7A7A")
        c.border = BOX
        if fmts and fmts[i]:
            c.number_format = fmts[i]


def total(ws, row, label, formulas, label_col=1):
    ws.cell(row=row, column=label_col, value=label).font = font(bold=True)
    for col, f in formulas.items():
        c = ws.cell(row=row, column=col, value=f)
        c.font = font(bold=True)
        c.number_format = MONEY
        c.fill = TOTAL_FILL
        c.border = BOX


def note(ws, cell, text):
    ws[cell] = text
    ws[cell].font = font(size=9, italic=True, color="5A6475")
    ws[cell].alignment = Alignment(wrap_text=True, vertical="top")


def protect(ws, landscape=False):
    # Print one page wide; inputs stay editable, labels and formulas are locked
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def build(sample=False):
    wb = Workbook()

    # ---------------- Start ----------------
    ws = wb.active
    ws.title = "Start"
    title(ws, "Falcon Tax Solutions · 2026 Tax Organizer", "Fill this out at home, then upload it to get your estimate.")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 96
    lines = [
        ("How it works", True),
        ("1. Fill in the YELLOW cells on each tab, left to right: Household, Jobs, Income, Expenses, Car, Trips, Payments, Checklist.", False),
        ("2. Grey rows show an example of what to enter. Leave them as they are.", False),
        ("3. Skip any tab that doesn't apply to you. Blank cells count as zero.", False),
        ("4. Save the file (Excel: .xlsx; Google Sheets: File > Download > Microsoft Excel).", False),
        ("5. Upload it on our Refund Estimator page. You'll see your estimated federal and state refund or amount due right away.", False),
        ("", False),
        ("Privacy", True),
        ("Do NOT enter Social Security numbers, bank account numbers or passwords in this file.", False),
        ("When you upload it, the file is read on your own device and is not sent to us. We collect IDs and documents only through our secure portal.", False),
        ("", False),
        ("Good to know", True),
        ("This is an estimate for the 2026 tax year (returns filed in 2027), not a filed return or a promise of any refund.", False),
        ("Married filing jointly: include both spouses' income. Married filing separately: include only your own.", False),
        ("Rideshare, delivery or any business car: fill in the Car tab (miles and car costs). We compare the IRS mileage rate with your actual costs and use the larger.", False),
        ("Truck drivers away from home overnight: list each trip on the Trips tab for the per diem deduction ($80/day, IRS Notice 2026-60).", False),
        ("Questions? Book a free consultation on our website.", False),
        ("", False),
        ("Color key", True),
    ]
    r = 4
    for text, bold in lines:
        c = ws.cell(row=r, column=2, value=text)
        c.font = font(size=12 if bold else 10, bold=bold, color=NAVY if bold else "1B2230")
        c.alignment = Alignment(wrap_text=True)
        r += 1
    for fill, text in [(INPUT_FILL, "Yellow: you fill this in"), (EXAMPLE_FILL, "Grey: example only"), (TOTAL_FILL, "Blue: calculated for you")]:
        ws.cell(row=r, column=1).fill = fill
        ws.cell(row=r, column=1).border = BOX
        ws.cell(row=r, column=2, value=text).font = font()
        r += 1
    protect(ws)

    # ---------------- Household ----------------
    hh = wb.create_sheet("Household")
    title(hh, "Household", "Filing status, state, and the people on your return.")
    for col, w in zip("ABCDE", [46, 26, 30, 18, 18]):
        hh.column_dimensions[col].width = w
    rows = [
        (4, "Filing status", "Choose from the list."),
        (5, "State you lived in during 2026", "Choose from the list."),
        (6, "Your birth year", "Example: 1985"),
        (7, "Spouse's birth year", "Married filing jointly only."),
        (8, "Did you pay for childcare so you could work?", "Yes or No"),
    ]
    for rr, lab, hint in rows:
        hh.cell(row=rr, column=1, value=lab).font = font(bold=True)
        inp(hh.cell(row=rr, column=2), YEAR if rr in (6, 7) else None)
        note(hh, f"C{rr}", hint)
    dv_status = DataValidation(type="list", formula1='"' + ",".join(s for s, _ in STATUSES) + '"', allow_blank=True)
    dv_state = DataValidation(type="list", formula1="=Data!$H$2:$H$52", allow_blank=True)
    dv_yes = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    for dv in (dv_status, dv_state, dv_yes):
        hh.add_data_validation(dv)
    dv_status.add("B4")
    dv_state.add("B5")
    dv_yes.add("B8")

    hh["A11"] = "Dependents (children and relatives you support)"
    hh["A11"].font = font(size=12, bold=True, color=NAVY)
    header(hh, 12, ["First name", "Relationship", "Birth year", "Months lived with you", "Has an SSN? (Yes/No)"])
    example(hh, 13, ["Example: Maya", "Daughter", 2016, 12, "Yes"], [None, None, YEAR, None, None])
    dv_yes2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, showErrorMessage=True, error='Choose Yes or No. Never type a Social Security number here.')
    hh.add_data_validation(dv_yes2)
    for rr in range(14, 24):
        for col in range(1, 6):
            inp(hh.cell(row=rr, column=col), YEAR if col == 3 else None)
        dv_yes2.add(f"E{rr}")
    note(hh, "A25", "Children under 17 at the end of 2026 were born in 2010 or later. Leave this table empty if you have no dependents.")
    hh.merge_cells("A25:E26")
    protect(hh)

    # ---------------- Jobs ----------------
    jb = wb.create_sheet("Jobs")
    title(jb, "Jobs (W-2 paychecks)", "One row per W-2. Use the final pay stub of the year if you don't have the W-2 yet.")
    header(jb, 4, ["Employer", "Whose job (You / Spouse)", "Wages (W-2 box 1)", "Federal tax withheld (box 2)",
                   "State tax withheld (box 17)", "Tips on W-2 (box 7)", "Overtime premium (extra half only)"],
           [28, 16, 16, 18, 18, 16, 20])
    jb.row_dimensions[4].height = 42
    example(jb, 5, ["Example: ABC Logistics", "You", 52000, 5200, 2400, 0, 1500], [None, None, MONEY, MONEY, MONEY, MONEY, MONEY])
    dv_who = DataValidation(type="list", formula1='"You,Spouse"', allow_blank=True)
    jb.add_data_validation(dv_who)
    for rr in range(6, 16):
        for col in range(1, 8):
            inp(jb.cell(row=rr, column=col), MONEY if col >= 3 else None)
        dv_who.add(f"B{rr}")
    total(jb, 16, "Total", {c: f"=SUM({chr(64 + c)}6:{chr(64 + c)}15)" for c in range(3, 8)})
    note(jb, "A18", "Overtime premium: only the extra 'half' of time-and-a-half (up to $12,500, or $25,000 joint, may be deductible). "
                    "Cash tips you did NOT report to your employer go on the Income tab.")
    jb.merge_cells("A18:G19")
    protect(jb, landscape=True)

    # ---------------- Income ----------------
    ic = wb.create_sheet("Income")
    title(ic, "Self-employment, cash & other income", "Business income, cash payments, and income not on a W-2.")
    ic["A4"] = "Self-employment & cash income (1099-NEC, 1099-K, cash, Zelle, checks)"
    ic["A4"].font = font(size=12, bold=True, color=NAVY)
    header(ic, 5, ["Payer or description", "Type", "Amount received"], [44, 18, 18])
    example(ic, 6, ["Example: Uber (1099-K)", "1099-K", 38000], [None, None, MONEY])
    dv_type = DataValidation(type="list", formula1='"1099-NEC,1099-K,Cash,Zelle / app,Check,Other"', allow_blank=True)
    ic.add_data_validation(dv_type)
    for rr in range(7, 19):
        for col in range(1, 4):
            inp(ic.cell(row=rr, column=col), MONEY if col == 3 else None)
        dv_type.add(f"B{rr}")
    total(ic, 19, "Total self-employment income", {3: "=SUM(C7:C18)"})
    note(ic, "A20", "Report all business income, including cash and app payments, even if you didn't get a 1099.")

    ic["A22"] = "Other income"
    ic["A22"].font = font(size=12, bold=True, color=NAVY)
    header(ic, 23, ["Type", "Form", "Amount"])
    other = [
        (24, "Cash tips NOT reported to your employer", "Tip log"),
        (25, "Bank interest", "1099-INT"),
        (26, "Unemployment benefits", "1099-G"),
        (27, "Retirement withdrawals (taxable amount)", "1099-R"),
        (28, "Other taxable income (prizes, rental profit, etc.)", ""),
    ]
    for rr, lab, form in other:
        ic.cell(row=rr, column=1, value=lab).font = font()
        ic.cell(row=rr, column=2, value=form).font = font(color="5A6475")
        inp(ic.cell(row=rr, column=3), MONEY)
    protect(ic)

    # ---------------- Expenses ----------------
    ex = wb.create_sheet("Expenses")
    title(ex, "Business expenses", "Self-employed only. Enter what you paid in 2026; the deductible amount is calculated.")
    header(ex, 4, ["Expense", "Schedule C line", "Amount paid", "Deductible share", "Deductible amount", "Notes"],
           [46, 12, 16, 12, 18, 46])
    example(ex, 5, ["Example: Fuel (diesel)", "9", 1250, 1, 1250, ""], [None, None, MONEY, "0%", MONEY, None])
    r = 6
    first = r
    for label, line, share, tip in EXPENSES:
        if line is None:
            c = ex.cell(row=r, column=1, value=label)
            c.font = font(bold=True, color=GOLD)
            r += 1
            continue
        if label:
            ex.cell(row=r, column=1, value=label).font = font()
        else:
            inp(ex.cell(row=r, column=1))
        ex.cell(row=r, column=2, value=line).font = font(color="5A6475")
        inp(ex.cell(row=r, column=3), MONEY)
        sc = ex.cell(row=r, column=4, value=share)
        sc.number_format = "0%"
        sc.font = font(color="5A6475")
        fc = ex.cell(row=r, column=5, value=f"=C{r}*D{r}")
        fc.number_format = MONEY
        if tip:
            note(ex, f"F{r}", tip)
        r += 1
    last = r - 1
    total(ex, r, "Total deductible business expenses", {3: f"=SUM(C{first}:C{last})", 5: f"=SUM(E{first}:E{last})"})
    exp_total_cell = f"E{r}"
    note(ex, f"A{r + 2}", "Equipment, electronics or software over $2,500 per item may need to be depreciated instead of fully expensed; we'll handle that on your return.")
    ex.merge_cells(f"A{r + 2}:F{r + 3}")
    protect(ex, landscape=True)

    # ---------------- Car ----------------
    car = wb.create_sheet("Car")
    title(car, "Business car (rideshare, delivery or any business car)", "Not for semi-trucks (use Expenses > Trucking). Keep your mileage log or the app's trip history.")
    car.column_dimensions["A"].width = 58
    car.column_dimensions["B"].width = 18
    car.column_dimensions["C"].width = 60
    header(car, 4, ["Miles", "Miles", "Notes"])
    for i, (key, lab) in enumerate(CAR_MILES):
        rr = 5 + i
        car.cell(row=rr, column=1, value=lab).font = font()
        inp(car.cell(row=rr, column=2), "#,##0")
    note(car, "C5", "Business miles: driving with the app on or for business (rideshare apps show 'online miles').")
    note(car, "C7", "Needed for the actual-cost method. Commuting and personal miles count here, not above.")
    header(car, 9, ["Car costs paid in 2026 (full amounts)", "Amount paid", "Notes"])
    first_cost = 10
    for i, (key, lab) in enumerate(CAR_COSTS):
        rr = first_cost + i
        car.cell(row=rr, column=1, value=lab).font = font()
        inp(car.cell(row=rr, column=2), MONEY)
    last_cost = first_cost + len(CAR_COSTS) - 1
    tc = last_cost + 1
    total(car, tc, "Total car costs", {2: f"=SUM(B{first_cost}:B{last_cost})"})
    r0 = tc + 2
    calc = [
        ("Standard mileage (72.5 cents Jan-Jun, 76 cents Jul-Dec)", f"=B5*Data!$C$66+B6*Data!$C$67", MONEY),
        ("Business share of all miles", '=IF(AND(B7>0,B5+B6>0),MIN(1,(B5+B6)/B7),"")', "0%"),
        ("Actual costs x business share", f'=IF(N(B{r0 + 1})>0,B{tc}*B{r0 + 1},"")', MONEY),
        ("We use the larger deduction", f'=MAX(B{r0},N(B{r0 + 2}))', MONEY),
    ]
    for i, (lab, f, fmt) in enumerate(calc):
        rr = r0 + i
        car.cell(row=rr, column=1, value=lab).font = font(bold=i == 3)
        c = car.cell(row=rr, column=2, value=f)
        c.number_format = fmt
        c.fill = TOTAL_FILL
        c.border = BOX
        c.font = font(bold=i == 3)
    note(car, f"A{r0 + 5}", "Tolls, parking, app fees and passenger supplies go on the Expenses tab (Rideshare & delivery); they count with either method. "
                            "If you used actual costs for this car in an earlier year, or a leased car started on one method, we'll confirm which method you must keep.")
    car.merge_cells(f"A{r0 + 5}:C{r0 + 7}")
    protect(car)
    car_rows = {"h1": "Car!B5", "h2": "Car!B6", "total": "Car!B7", "costs": f"Car!B{tc}"}

    # ---------------- Trips ----------------
    tp = wb.create_sheet("Trips")
    title(tp, "Truck driver trips (per diem)", "Self-employed drivers away from home overnight. One row per trip.")
    header(tp, 4, ["Left home (date)", "Got back (date)", "Full days in Canada", "Full days away", "Travel days (75%)", "Per diem", "Check"],
           [18, 18, 16, 16, 16, 16, 24])
    example(tp, 5, ["2026-03-02", "2026-03-09", 0, 6, 2, 600], [None, None, None, None, None, MONEY])
    for rr in range(6, 36):
        inp(tp.cell(row=rr, column=1), "yyyy-mm-dd")
        inp(tp.cell(row=rr, column=2), "yyyy-mm-dd")
        inp(tp.cell(row=rr, column=3), "0")
        ok = f"AND(ISNUMBER(A{rr}),ISNUMBER(B{rr}),B{rr}>A{rr})"
        in_year = lambda c: f"AND({c}>=DATE(2026,1,1),{c}<=DATE(2026,12,31))"
        # only 2026 days count
        tp.cell(row=rr, column=4, value=f'=IF({ok},MAX(0,MIN(B{rr}-1,DATE(2026,12,31))-MAX(A{rr}+1,DATE(2026,1,1))+1),0)').number_format = "0"
        tp.cell(row=rr, column=5, value=f'=IF({ok},{in_year(f"A{rr}")}*1+{in_year(f"B{rr}")}*1,0)').number_format = "0"
        f = tp.cell(row=rr, column=6, value=f"=(D{rr}-MIN(C{rr},D{rr}))*Data!$C$60+MIN(C{rr},D{rr})*Data!$C$61+E{rr}*Data!$C$60*Data!$C$62")
        f.number_format = MONEY
        # warn when two trips share days (the website counts each day once)
        tp.cell(row=rr, column=7, value=f'=IF(AND({ok},SUMPRODUCT(ISNUMBER($A$6:$A$35)*ISNUMBER($B$6:$B$35)*($A$6:$A$35<B{rr})*($B$6:$B$35>A{rr}))>1),"Overlaps another trip","")').font = font(color="A3472B", bold=True)
    tr = 36
    tp.cell(row=tr, column=1, value="Total").font = font(bold=True)
    for col, f in [(3, "=SUM(C6:C35)"), (4, "=SUM(D6:D35)"), (5, "=SUM(E6:E35)"), (6, "=SUM(F6:F35)")]:
        c = tp.cell(row=tr, column=col, value=f)
        c.font = font(bold=True)
        c.fill = TOTAL_FILL
        c.border = BOX
        c.number_format = MONEY if col == 6 else "0"
    tp.cell(row=37, column=5, value="Deductible (80%)").font = font(bold=True)
    c = tp.cell(row=37, column=6, value="=F36*Data!$C$63")
    c.number_format = MONEY
    c.font = font(bold=True)
    c.fill = TOTAL_FILL
    note(tp, "A39", "IRS transportation per diem (Notice 2026-60): $80/day in the U.S., $86 in Canada. The day you leave and the day you "
                    "return count at 75%. Only 2026 days count; don't list the same day on two trips. Drivers under DOT hours-of-service rules deduct 80%. Employee (W-2) drivers can't deduct this on "
                    "their federal return. Keep your logbook or ELD records.")
    tp.merge_cells("A39:F41")
    protect(tp)

    # ---------------- Payments ----------------
    pm = wb.create_sheet("Payments")
    title(pm, "Estimated tax payments", "Quarterly payments you sent to the IRS or your state for 2026.")
    header(pm, 4, ["Quarter", "Federal paid", "Date paid", "State paid", "Date paid"], [22, 16, 16, 16, 16])
    for i, (q, due) in enumerate([("Q1 (due Apr 15, 2026)", 0), ("Q2 (due Jun 15, 2026)", 0), ("Q3 (due Sep 15, 2026)", 0), ("Q4 (due Jan 15, 2027)", 0)]):
        rr = 5 + i
        pm.cell(row=rr, column=1, value=q).font = font()
        for col, fmt in [(2, MONEY), (3, "yyyy-mm-dd"), (4, MONEY), (5, "yyyy-mm-dd")]:
            inp(pm.cell(row=rr, column=col), fmt)
    total(pm, 9, "Total", {2: "=SUM(B5:B8)", 4: "=SUM(D5:D8)"})
    protect(pm)

    # ---------------- Checklist ----------------
    ck = wb.create_sheet("Checklist")
    title(ck, "Checklist", "Answer Yes or No. These help us find credits and documents.")
    ck.column_dimensions["A"].width = 76
    ck.column_dimensions["B"].width = 12
    dv_yes3 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ck.add_data_validation(dv_yes3)
    for i, (_, q) in enumerate(CHECKLIST):
        rr = 4 + i
        ck.cell(row=rr, column=1, value=q).font = font()
        inp(ck.cell(row=rr, column=2))
        dv_yes3.add(f"B{rr}")
    protect(ck)

    # ---------------- Data (read by the website) ----------------
    dt = wb.create_sheet("Data")
    title(dt, "Data for the estimator", "Calculated from your answers. Please don't change this tab.")
    dt["A1"].value = "Data for the estimator"
    for col, w in zip("ABCDEFGHI", [22, 52, 16, 2, 26, 8, 2, 22, 6]):
        dt.column_dimensions[col].width = w
    header_cells = [("A3", "Key"), ("B3", "Meaning"), ("C3", "Value")]
    for ref, text in header_cells:
        dt[ref] = text
        dt[ref].font = font(bold=True, color="FFFFFF")
        dt[ref].fill = HEAD_FILL
    # lookup tables
    dt["E1"], dt["F1"] = "Filing status", "Code"
    for i, (name, code) in enumerate(STATUSES):
        dt.cell(row=2 + i, column=5, value=name)
        dt.cell(row=2 + i, column=6, value=code)
    dt["H1"], dt["I1"] = "State", "Code"
    for i, (code, name) in enumerate(sorted(STATES.items(), key=lambda kv: kv[1])):
        dt.cell(row=2 + i, column=8, value=name)
        dt.cell(row=2 + i, column=9, value=code)

    yes = lambda ref: f'=IF(TRIM({ref})="Yes",1,0)'
    keys = [
        ("format", "Organizer version", '="falcon-organizer-2026"'),
        ("status", "Filing status code", '=IFERROR(INDEX($F$2:$F$5,MATCH(Household!B4,$E$2:$E$5,0)),"single")'),
        ("state", "State code", '=IFERROR(INDEX($I$2:$I$52,MATCH(Household!B5,$H$2:$H$52,0)),"")'),
        ("you65", "You are 65 or older (born 1961 or earlier)", '=IF(AND(ISNUMBER(Household!B6),Household!B6>0,Household!B6<=1961),1,0)'),
        ("spouse65", "Spouse is 65 or older (joint returns)", '=IF(AND(C5="mfj",ISNUMBER(Household!B7),Household!B7>0,Household!B7<=1961),1,0)'),
        ("childrenUnder17", "Children under 17 (born 2010 or later)", '=COUNTIFS(Household!A14:A23,"?*",Household!C14:C23,">=2010")'),
        ("otherDependents", "Other dependents", '=COUNTIF(Household!A14:A23,"?*")-C9'),
        ("hasDeps", "Has dependents", '=IF(C9+C10>0,"yes","no")'),
        ("childcare", "Paid childcare", yes("Household!B8")),
        ("wages", "W-2 wages plus unreported cash tips", "=Jobs!C16+Income!C24"),
        ("withheld", "Federal tax withheld", "=Jobs!D16"),
        ("stateWithheld", "State tax withheld", "=Jobs!E16"),
        ("tips", "Tips (W-2 and cash)", "=Jobs!F16+Income!C24"),
        ("overtimePremium", "Overtime premium", "=Jobs!G16"),
        ("seIncome", "Self-employment income", "=Income!C19"),
        ("seExpenses", "Deductible business expenses", f"=Expenses!{exp_total_cell}"),
        ("estimatedPayments", "Federal estimated payments", "=Payments!B9"),
        ("stateEstimated", "State estimated payments", "=Payments!D9"),
        ("isTrucker", "Logged truck driver trips", "=IF(Trips!D36+Trips!E36>0,1,0)"),
        ("perDiemFull", "Full days away (U.S. and Canada)", "=MIN(365,Trips!D36)"),
        ("perDiemCanada", "Of those, full days in Canada", "=MIN(C23,SUMPRODUCT((Trips!C6:C35<Trips!D6:D35)*Trips!C6:C35+(Trips!C6:C35>=Trips!D6:D35)*Trips!D6:D35))"),
        ("perDiemPartial", "Travel days (75%)", "=Trips!E36"),
        ("usesCar", "Uses a car for business (Car tab)", f'=IF(N({car_rows["h1"]})+N({car_rows["h2"]})+N({car_rows["costs"]})>0,1,0)'),
        ("carMilesH1", "Business miles Jan-Jun", f'=N({car_rows["h1"]})'),
        ("carMilesH2", "Business miles Jul-Dec", f'=N({car_rows["h2"]})'),
        ("carMilesTotal", "All miles driven in 2026", f'=N({car_rows["total"]})'),
        ("carCosts", "Actual car costs (before business share)", f'=N({car_rows["costs"]})'),
        ("interest", "Bank interest", "=Income!C25"),
        ("unemployment", "Unemployment", "=Income!C26"),
        ("otherIncome", "Retirement withdrawals and other income", "=Income!C27+Income!C28"),
        ("retirement", "Has retirement withdrawals", '=IF(Income!C27>0,1,0)'),
    ]
    for i, (key, _) in enumerate(CHECKLIST):
        keys.append((key, CHECKLIST[i][1], yes(f"Checklist!B{4 + i}")))
    for i, (k, meaning, f) in enumerate(keys):
        rr = 4 + i
        dt.cell(row=rr, column=1, value=k).font = font(bold=True)
        dt.cell(row=rr, column=2, value=meaning).font = font(color="5A6475")
        dt.cell(row=rr, column=3, value=f).font = font()
    # per diem rates referenced by the Trips tab (fixed rows)
    assert 4 + len(keys) < 59, "Data keys overlap the rate rows"
    dt["A59"] = "Per diem rates (IRS Notice 2026-60)"
    dt["A59"].font = font(bold=True, color=NAVY)
    dt["A65"] = "Standard mileage rates 2026 (Notice 2026-10, Announcement 2026-11)"
    dt["A65"].font = font(bold=True, color=NAVY)
    for rr, lab, val in [(66, "Jan 1 - Jun 30, per mile", MILEAGE_RATES[0]), (67, "Jul 1 - Dec 31, per mile", MILEAGE_RATES[1])]:
        dt.cell(row=rr, column=2, value=lab).font = font()
        c = dt.cell(row=rr, column=3, value=val)
        c.number_format = '$0.000'
        c.font = font(color="0000FF")
    for rr, lab, val, fmt in [(60, "U.S. daily rate", 80, MONEY), (61, "Canada daily rate", 86, MONEY),
                              (62, "Travel day share", 0.75, "0%"), (63, "Deductible share (DOT hours)", 0.8, "0%")]:
        dt.cell(row=rr, column=2, value=lab).font = font()
        c = dt.cell(row=rr, column=3, value=val)
        c.number_format = fmt
        c.font = font(color="0000FF")
    dt["C60"].comment = Comment("Source: IRS Notice 2026-60, special M&IE rate for the transportation industry, Oct 1, 2026 - Sep 30, 2027 (same as prior year).", "Falcon Tax")
    protect(dt, landscape=True)

    if sample:
        fill_sample(wb)
    return wb


def fill_sample(wb):
    """Realistic answers used to test the website import (married trucker in Illinois)."""
    import datetime as d
    hh = wb["Household"]
    hh["B4"], hh["B5"], hh["B6"], hh["B7"], hh["B8"] = "Married filing jointly", "Illinois", 1984, 1987, "No"
    for r, row in enumerate([("Maya", "Daughter", 2016, 12, "Yes"), ("Leo", "Son", 2019, 12, "Yes"), ("Rosa", "Mother", 1958, 12, "Yes")]):
        for c, v in enumerate(row):
            hh.cell(row=14 + r, column=1 + c, value=v)
    jb = wb["Jobs"]
    for c, v in enumerate(["County School District", "Spouse", 41000, 3100, 1900, 0, 0]):
        jb.cell(row=6, column=1 + c, value=v)
    ic = wb["Income"]
    for c, v in enumerate(["Freight broker loads (1099-NEC)", "1099-NEC", 148000]):
        ic.cell(row=7, column=1 + c, value=v)
    for c, v in enumerate(["Cash and Zelle loads", "Zelle / app", 6500]):
        ic.cell(row=8, column=1 + c, value=v)
    ic["C25"] = 120
    ex = wb["Expenses"]
    amounts = {"Fuel (diesel)": 52000, "IFTA fuel tax paid": 1300, "Form 2290 heavy vehicle use tax": 550,
               "UCR registration": 46, "IRP plates & registration": 1800, "Truck insurance (liability, cargo, physical damage)": 14000,
               "Load boards (DAT, Truckstop, 123Loadboard)": 1980, "ELD subscription": 420, "Tolls, scales & parking": 2100,
               "Brakes": 1240, "Tires (new or retread)": 2860, "Oil & filter change": 780, "Truck wash": 300,
               "Business meals (enter full amount)": 0}
    for r in range(6, ex.max_row + 1):
        lab = ex.cell(row=r, column=1).value
        if lab in amounts:
            ex.cell(row=r, column=3, value=amounts[lab])
    tp = wb["Trips"]
    start = d.date(2026, 1, 5)
    for i in range(20):
        leave = start + d.timedelta(days=i * 16)
        tp.cell(row=6 + i, column=1, value=leave)
        tp.cell(row=6 + i, column=2, value=leave + d.timedelta(days=10))
        tp.cell(row=6 + i, column=3, value=2 if i % 5 == 0 else 0)
    pm = wb["Payments"]
    for i in range(4):
        pm.cell(row=5 + i, column=2, value=4000)
        pm.cell(row=5 + i, column=4, value=900)
    ck = wb["Checklist"]
    for i in range(len(CHECKLIST)):
        ck.cell(row=4 + i, column=2, value="Yes" if CHECKLIST[i][0] == "mortgage" else "No")


if __name__ == "__main__":
    out = next((a for a in sys.argv[1:] if not a.startswith("--")), "falcon-tax-organizer-2026.xlsx")
    build(sample="--sample" in sys.argv).save(out)
    print("wrote", out)
